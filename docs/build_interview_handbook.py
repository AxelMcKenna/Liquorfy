#!/usr/bin/env python3
"""Build the Liquorfy interview handbook DOCX from its Markdown source."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "LIQUORFY_INTERVIEW_HANDBOOK.md"
OUTPUT = ROOT / "Liquorfy-Interview-Handbook.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "243240"
MUTED = "5D6B78"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF4D6"
PALE_GREY = "F4F6F8"
WHITE = "FFFFFF"
TABLE_WIDTH = 9360

PAGE_BREAK_SECTIONS = {
    "3. Architecture",
    "6. Backend internals",
    "8. The scraper subsystem",
    "11. The React web application",
    "13. Authentication, security, and privacy",
    "15. Testing strategy and current coverage",
    "17. Technical debt and the best improvement roadmap",
    "18. Interview question bank",
    "22. Rapid-review cheat sheet",
}

GUIDE_MAP = (
    ("Story", "1–2", "Pitch, domain, and product judgement"),
    ("System", "3–7", "Architecture, flows, API, search, and data"),
    ("Collection", "8–10", "Scrapers, normalisation, and freshness"),
    ("Clients", "11–12", "React web and native SwiftUI"),
    ("Operations", "13–17", "Security, deployment, testing, and debt"),
    ("Interview", "18–23", "Q&A, design drills, stories, and revision"),
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH:
        raise ValueError(f"table widths must sum to {TABLE_WIDTH}: {widths}")

    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for index, cell in enumerate(row.cells):
            width = widths[index]
            cell.width = Twips(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def choose_widths(column_count: int) -> list[int]:
    presets = {
        1: [9360],
        2: [2850, 6510],
        3: [1700, 2350, 5310],
        4: [1700, 2100, 2300, 3260],
        5: [1450, 1700, 1750, 2000, 2460],
    }
    if column_count in presets:
        return presets[column_count]
    base = TABLE_WIDTH // column_count
    widths = [base] * column_count
    widths[-1] += TABLE_WIDTH - sum(widths)
    return widths


def set_run_font(run, name: str, size: float | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, *, base_size: float | None = None) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            if base_size:
                run.font.size = Pt(base_size)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Consolas", 8.5 if base_size is None else min(base_size, 8.5), DARK_BLUE)
            run.font.highlight_color = None
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if base_size:
                run.font.size = Pt(base_size)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            hyperlink = OxmlElement("w:hyperlink")
            rel_id = paragraph.part.relate_to(
                url,
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            hyperlink.set(qn("r:id"), rel_id)
            new_run = OxmlElement("w:r")
            r_pr = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), BLUE)
            underline = OxmlElement("w:u")
            underline.set(qn("w:val"), "single")
            r_pr.extend((color, underline))
            text_node = OxmlElement("w:t")
            text_node.text = label
            new_run.extend((r_pr, text_node))
            hyperlink.append(new_run)
            paragraph._p.append(hyperlink)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        if base_size:
            run.font.size = Pt(base_size)


def set_paragraph_keep(paragraph, *, keep_next=False, keep_lines=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    if keep_lines:
        node = OxmlElement("w:keepLines")
        p_pr.append(node)
    widow = OxmlElement("w:widowControl")
    p_pr.append(widow)


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    headings = (
        ("Title", 30, DARK_BLUE, 0, 12),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    )
    for name, size, color, before, after in headings:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string(INK)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.1)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.05

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Callout"]
    callout.font.name = "Calibri"
    callout.font.size = Pt(10.5)
    callout.font.italic = True
    callout.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.08)
    callout.paragraph_format.space_before = Pt(7)
    callout.paragraph_format.space_after = Pt(9)
    callout.paragraph_format.line_spacing = 1.15


def add_paragraph_shading(paragraph, fill: str, *, border_color: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if border_color:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border_color)
        borders.append(left)
        p_pr.append(borders)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PAGE ")
    set_run_font(run, "Calibri", 8, MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def new_numbering_id(doc: Document, start: int) -> int:
    """Create a real numbered-list instance with an explicit starting value."""

    numbering = doc.part.numbering_part.element
    base_num_id = int(doc.styles["List Number"]._element.pPr.numPr.numId.val)
    base_num = next(
        node
        for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == base_num_id
    )
    abstract_num_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_ids, default=0) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_num_id)
    num.append(abstract)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), "0")
    num_id_node = num_pr.find(qn("w:numId"))
    if num_id_node is None:
        num_id_node = OxmlElement("w:numId")
        num_pr.append(num_id_node)
    num_id_node.set(qn("w:val"), str(num_id))


def add_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("LIQUORFY  /  INTERVIEW HANDBOOK")
    set_run_font(run, "Calibri", 8, MUTED)
    run.bold = True

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    add_page_field(p)


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = kicker.add_run("INTERVIEW PREPARATION GUIDE")
    set_run_font(run, "Calibri", 10, BLUE)
    run.bold = True

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(10)
    title.add_run("Liquorfy")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(22)
    run = subtitle.add_run("The code-grounded handbook for explaining the product, architecture, trade-offs, and next steps")
    set_run_font(run, "Calibri", 16, INK)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(18)
    run = rule.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    set_run_font(run, "Calibri", 11, BLUE)

    meta = doc.add_paragraph()
    add_inline(meta, "**Repository snapshot:** `main` at `c16e42a` (21 June 2026)")
    meta = doc.add_paragraph()
    add_inline(meta, "**Inspection date:** 23 July 2026")
    meta = doc.add_paragraph()
    add_inline(meta, "**Focus:** full-stack, backend, mobile, platform, and system-design interviews")

    tag = doc.add_paragraph()
    tag.paragraph_format.space_before = Pt(22)
    add_paragraph_shading(tag, PALE_BLUE, border_color=BLUE)
    run = tag.add_run("CODE-GROUNDED  •  CANDID  •  CURRENT")
    set_run_font(run, "Calibri", 10, DARK_BLUE)
    run.bold = True

    doc.add_paragraph()
    note = doc.add_paragraph()
    add_inline(
        note,
        "Scope note: this guide treats the committed revision as the interview baseline and labels newer working-tree hardening separately. Older documentation is not treated as authoritative when it conflicts with code.",
        base_size=9.5,
    )
    note.paragraph_format.space_before = Pt(32)
    note.paragraph_format.space_after = Pt(0)

    doc.add_page_break()


def add_guide_map(doc: Document) -> None:
    heading = doc.add_paragraph("Guide map", style="Heading 1")
    set_paragraph_keep(heading, keep_next=True)
    intro = doc.add_paragraph(
        "Use the first two columns to choose a revision route; use the final column to remember what each block proves in an interview."
    )
    intro.paragraph_format.space_after = Pt(9)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for idx, text in enumerate(("Route", "Sections", "What you should be able to defend")):
        cell = table.rows[0].cells[idx]
        cell.text = text
        set_cell_shading(cell, PALE_BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            set_run_font(run, "Calibri", 9.5, DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for route, sections, description in GUIDE_MAP:
        cells = table.add_row().cells
        for idx, text in enumerate((route, sections, description)):
            cells[idx].text = text
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    set_run_font(run, "Calibri", 9.3, INK)
    set_table_geometry(table, [1550, 1450, 6360])
    doc.add_paragraph()


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    header = rows[0]
    data = rows[2:] if all(re.fullmatch(r"\s*:?-{3,}:?\s*", item) for item in rows[1]) else rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"

    for idx, text in enumerate(header):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        add_inline(cell.paragraphs[0], text.strip(), base_size=9.2)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        set_cell_shading(cell, PALE_BLUE)
    set_repeat_table_header(table.rows[0])

    for row_idx, row in enumerate(data):
        values = row + [""] * (len(header) - len(row))
        cells = table.add_row().cells
        for idx, text in enumerate(values[: len(header)]):
            cell = cells[idx]
            cell.text = ""
            add_inline(cell.paragraphs[0], text.strip(), base_size=9.0)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.03
            if row_idx % 2:
                set_cell_shading(cell, "F8FAFC")
    set_table_geometry(table, choose_widths(len(header)))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph(style="Code Block")
    add_paragraph_shading(paragraph, PALE_GREY, border_color="B8C4CF")
    for idx, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run, "Consolas", 8.5, INK)
        if idx != len(lines) - 1:
            run.add_break()
    set_paragraph_keep(paragraph)


def markdown_rows(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        raw = lines[index].strip().strip("|")
        rows.append([cell.strip() for cell in raw.split("|")])
        index += 1
    return rows, index


def add_body_from_markdown(doc: Document, text: str) -> None:
    lines = text.splitlines()
    start = next(idx for idx, line in enumerate(lines) if line.startswith("## How to use"))
    index = start
    paragraph_buffer: list[str] = []
    active_numbering_id: int | None = None

    def flush_buffer() -> None:
        nonlocal paragraph_buffer, active_numbering_id
        if not paragraph_buffer:
            return
        paragraph = doc.add_paragraph()
        add_inline(paragraph, " ".join(piece.strip() for piece in paragraph_buffer))
        set_paragraph_keep(paragraph)
        paragraph_buffer = []
        active_numbering_id = None

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_buffer()
            active_numbering_id = None
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(doc, code_lines)
            index += 1
            continue

        if stripped.startswith("|"):
            flush_buffer()
            active_numbering_id = None
            rows, index = markdown_rows(lines, index)
            add_markdown_table(doc, rows)
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            flush_buffer()
            active_numbering_id = None
            hashes, title = heading_match.groups()
            level = min(len(hashes) - 1, 3)
            if level == 1 and title in PAGE_BREAK_SECTIONS and len(doc.paragraphs) > 2:
                doc.add_page_break()
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            add_inline(paragraph, title)
            set_paragraph_keep(paragraph, keep_next=True)
            index += 1
            continue

        bullet_match = re.match(r"^\s*[-*]\s+(.+)$", line)
        number_match = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        if bullet_match or number_match:
            flush_buffer()
            style = "List Bullet" if bullet_match else "List Number"
            content = bullet_match.group(1) if bullet_match else number_match.group(2)
            paragraph = doc.add_paragraph(style=style)
            add_inline(paragraph, content)
            if number_match:
                if active_numbering_id is None:
                    active_numbering_id = new_numbering_id(doc, int(number_match.group(1)))
                apply_numbering(paragraph, active_numbering_id)
            else:
                active_numbering_id = None
            set_paragraph_keep(paragraph)
            index += 1
            continue

        if stripped.startswith(">"):
            flush_buffer()
            active_numbering_id = None
            paragraph = doc.add_paragraph(style="Callout")
            add_inline(paragraph, stripped.lstrip("> ").strip())
            add_paragraph_shading(paragraph, PALE_GOLD, border_color="D49B22")
            set_paragraph_keep(paragraph)
            index += 1
            continue

        if not stripped:
            flush_buffer()
            active_numbering_id = None
            index += 1
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_buffer()


def set_document_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True
    add_header_footer(section)

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    props = doc.core_properties
    props.title = "Liquorfy Interview Handbook"
    props.subject = "Code-grounded interview preparation for the Liquorfy project"
    props.author = "Liquorfy project team"
    props.keywords = "Liquorfy, interview, FastAPI, React, SwiftUI, PostgreSQL, PostGIS, scraping"
    props.comments = "Generated from repository inspection on 23 July 2026."


def main() -> int:
    if not SOURCE.exists():
        print(f"missing source: {SOURCE}", file=sys.stderr)
        return 1

    doc = Document()
    configure_styles(doc)
    set_document_defaults(doc)
    add_cover(doc)
    add_guide_map(doc)
    add_body_from_markdown(doc, SOURCE.read_text(encoding="utf-8"))
    doc.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
